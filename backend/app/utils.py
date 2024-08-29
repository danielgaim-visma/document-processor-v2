import os
import re
import docx
import openpyxl
import nltk
from nltk.tokenize import word_tokenize
from nltk.corpus import stopwords
import openai
import json
from typing import List, Dict, Optional
import PyPDF2
import logging
from logging_config import log_function_call, log_error

logger = logging.getLogger('file_processing')

nltk.download('punkt', quiet=True)
nltk.download('stopwords', quiet=True)

openai.api_key = os.getenv("OPENAI_API_KEY")

def read_file_content(file_path: str) -> Optional[str]:
    log_function_call(logger, 'read_file_content', file_path=file_path)
    file_extension = os.path.splitext(file_path)[1].lower()

    if not os.path.exists(file_path):
        log_error(logger, f"File not found: {file_path}")
        return None

    try:
        if file_extension == '.txt':
            with open(file_path, 'r', encoding='utf-8') as file:
                return file.read()
        elif file_extension == '.docx':
            doc = docx.Document(file_path)
            return ' '.join([para.text for para in doc.paragraphs])
        elif file_extension == '.xlsx':
            workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            sheet = workbook.active
            return '\n'.join([' '.join([str(cell.value) if cell.value is not None else '' for cell in row]) for row in sheet.iter_rows()])
        elif file_extension == '.pdf':
            with open(file_path, 'rb') as file:
                reader = PyPDF2.PdfReader(file)
                return ' '.join([page.extract_text() for page in reader.pages])
        else:
            log_error(logger, f"Unsupported file type: {file_extension}")
            return None
    except Exception as e:
        log_error(logger, f"Error reading file {file_path}: {str(e)}", exc_info=True)
        return None

def parse_content(content):
    log_function_call(logger, 'parse_content', content_length=len(content))
    try:
        sections = re.split(r'\n(?=#{1,3}\s)', content)  # Split on headings
        if len(sections) <= 1:
            words = content.split()
            sections = [' '.join(words[i:i + 500]) for i in range(0, len(words), 500)]
        logger.info(f"Parsed content into {len(sections)} sections")
        return sections
    except Exception as e:
        log_error(logger, f"Error parsing content: {str(e)}", exc_info=True)
        raise

def extract_keywords(text):
    log_function_call(logger, 'extract_keywords', text_length=len(text))
    try:
        stop_words = set(stopwords.words('norwegian'))
        words = word_tokenize(text.lower())
        filtered_words = [word for word in words if word.isalnum() and word not in stop_words]
        word_freq = nltk.FreqDist(filtered_words)
        keywords = [word for word, _ in word_freq.most_common(6)]
        logger.info(f"Extracted keywords: {keywords}")
        return keywords
    except Exception as e:
        log_error(logger, f"Error extracting keywords: {str(e)}", exc_info=True)
        raise

def extract_and_structure_data(content: str, keywords: List[str]) -> Dict:
    log_function_call(logger, 'extract_and_structure_data', content_length=len(content), keywords=keywords)
    system_instruction = {
        "role": "system",
        "content": """Du er en hjelpsom assistent som ekstraherer, strukturerer og organiserer informasjon fra tekst. Din oppgave er å:
        1. Identifisere og kategorisere nøkkelbegreper, enheter (f.eks. navn, datoer, steder) og relasjoner innen teksten.
        2. Strukturere informasjonen i et klart og logisk format, for eksempel JSON, med riktige etiketter og hierarki.
        3. Sørge for at resultatet er både lesbart for mennesker og optimalisert for søkbarhet av en LLM, inkludert indekserte nøkkelord og metadata-tagging.
        4. Tilpasse ekstraheringen basert på konteksten.
        5. Prioritere klarhet, konsistens og fullstendighet i strukturen."""
    }

    user_prompt = f"""Analyser følgende innhold og nøkkelord, og ekstraher og strukturer hovedpunktene:

    Innhold: {content}

    Nøkkelord: {', '.join(keywords)}

    Vennligst gi et strukturert resultat i JSON-format med følgende felter:
    - title: Tittel på seksjonen (hvis ingen tydelig tittel, bruk 'Untitled Section')
    - body: Hovedinnholdet i seksjonen
    - tags: Liste over relevante nøkkelord fra innholdet"""

    try:
        logger.info(f"Sending prompt to OpenAI. Content length: {len(content)}")
        response = openai.ChatCompletion.create(
            model="gpt-4o-2024-08-06",
            messages=[
                system_instruction,
                {"role": "user", "content": user_prompt}
            ],
            temperature=0.4,
            max_tokens=4096
        )

        if not response.choices or not response.choices[0].message['content'].strip():
            log_error(logger, "OpenAI API returned an empty response")
            return {
                "title": "Error: Empty Response",
                "body": "The AI model returned an empty response for this section.",
                "tags": []
            }

        result = response.choices[0].message['content'].strip()
        logger.info(f"Received response from OpenAI. Response length: {len(result)}")

        # Check if the response starts with ```json and ends with ```
        if result.startswith("```json") and result.endswith("```"):
            # Remove the ```json and ``` markers
            result = result[7:-3].strip()

        # Parse the JSON result
        try:
            parsed_result = json.loads(result)
            return parsed_result
        except json.JSONDecodeError as json_error:
            log_error(logger, f"Failed to parse JSON from OpenAI response. Response content: {result}")
            return {
                "title": "Error: Invalid JSON",
                "body": f"Failed to parse JSON from OpenAI response: {result}",
                "tags": []
            }

    except openai.error.OpenAIError as openai_error:
        log_error(logger, f"OpenAI API error: {str(openai_error)}")
        return {
            "title": "Error: OpenAI API",
            "body": f"OpenAI API error: {str(openai_error)}",
            "tags": []
        }
    except Exception as e:
        log_error(logger, f"Unexpected error in OpenAI API call: {str(e)}", exc_info=True)
        return {
            "title": "Error: Unexpected",
            "body": f"Unexpected error in OpenAI API call: {str(e)}",
            "tags": []
        }
def process_file(file_path: str) -> str:
    content = read_file_content(file_path)
    keywords = extract_keywords(content)
    sections = parse_content(content)

    output_dir = "processed_files"
    os.makedirs(output_dir, exist_ok=True)

    json_files = []
    for i, section in enumerate(sections):
        structured_data = extract_and_structure_data(section, keywords)
        file_name = f"section_{i+1}.json"
        file_path = os.path.join(output_dir, file_name)
        with open(file_path, 'w', encoding='utf-8') as f:
            json.dump(structured_data, f, ensure_ascii=False, indent=2)
        json_files.append(file_path)

    zip_file_name = f"processed_files_{datetime.now().strftime('%Y%m%d_%H%M%S')}.zip"
    with zipfile.ZipFile(zip_file_name, 'w') as zipf:
        for file in json_files:
            zipf.write(file, arcname=os.path.basename(file))

    return zip_file_name
if __name__ == "__main__":
    file_path = "path_to_your_file.txt"  # Replace with the actual file path
    zip_file = process_file(file_path)
    print(f"Processing complete. Output saved in {zip_file}")